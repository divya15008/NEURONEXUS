"""Extended API coverage: validation, error paths, library, decisions, AI status."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from backend.main import app

client = TestClient(app)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_docx(path: Path, paragraphs: int = 3) -> None:
    doc = Document()
    doc.add_heading("Coverage Test Manuscript", level=1)
    doc.add_paragraph("Test Author")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("An abstract describing the study and its method. [1]")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Background and motivation.")
    for _ in range(paragraphs):
        doc.add_paragraph("Supporting paragraph.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Author T. Test study. Journal 2024.")
    doc.save(path)


def upload_sample(tmp_path: Path, name: str = "sample.docx") -> str:
    sample = tmp_path / name
    make_docx(sample)
    with sample.open("rb") as fh:
        response = client.post("/api/documents/upload", files={"file": (name, fh, DOCX)})
    assert response.status_code == 200, response.text
    return response.json()["document_id"]


def test_health():
    assert client.get("/health").json() == {"status": "ok"}


def test_ai_status_has_no_secrets():
    body = client.get("/api/ai/status").json()
    assert body["mode"] in {"local", "groq", "grok", "gemini"}
    assert "key" not in body and "KEY" not in str(body)


def test_upload_rejects_non_docx(tmp_path):
    response = client.post(
        "/api/documents/upload",
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )
    assert response.status_code == 400


def test_upload_rejects_oversized(tmp_path, monkeypatch):
    import backend.api.documents as documents_api

    monkeypatch.setattr(documents_api, "MAX_UPLOAD_BYTES", 10)
    sample = tmp_path / "big.docx"
    make_docx(sample)
    with sample.open("rb") as fh:
        response = client.post("/api/documents/upload", files={"file": ("big.docx", fh, DOCX)})
    assert response.status_code == 413


def test_get_missing_document_returns_404():
    assert client.get("/api/documents/DOC-NOPE").status_code == 404


def test_analyze_missing_document_returns_404():
    assert client.post("/api/documents/DOC-NOPE/analyze", json={}).status_code == 404


def test_structure_requires_analysis(tmp_path):
    doc_id = upload_sample(tmp_path, "unanalyzed.docx")
    assert client.get(f"/api/documents/{doc_id}/structure").status_code == 404


def test_documents_library_lists_uploads(tmp_path):
    doc_id = upload_sample(tmp_path, "listed.docx")
    body = client.get("/api/documents").json()
    ids = [item["document_id"] for item in body["documents"]]
    assert doc_id in ids
    entry = next(item for item in body["documents"] if item["document_id"] == doc_id)
    assert entry["filename"].endswith(".docx")
    assert entry["has_analysis"] is False
    assert entry["has_generated"] is False


def test_citation_check_with_empty_inputs():
    body = client.post("/api/citations/check", json={"citations": [], "references": []}).json()
    assert body["missing_references"] == []
    assert body["uncited_references"] == []


def test_citation_flags_missing_and_uncited():
    body = client.post(
        "/api/citations/check",
        json={
            "citations": [{"text": "[9]", "citation_type": "numeric", "reference_ids": ["9"]}],
            "references": [{"id": "9", "raw_text": "Ninth entry."}, {"id": "10", "raw_text": "Never cited."}],
        },
    ).json()
    assert body["missing_references"] == []
    assert len(body["uncited_references"]) == 1


def test_journal_rules_unknown_returns_404():
    assert client.get("/api/journals/unknown-journal/rules").status_code == 404


def test_all_analysis_modules_respond(tmp_path):
    doc_id = upload_sample(tmp_path, "modules.docx")
    analysis = client.post(f"/api/documents/{doc_id}/analyze", json={}).json()
    for module in ["quality", "novelty", "methodology", "contribution", "writing"]:
        response = client.post(f"/api/{module}/analyze", json={"document_id": doc_id, "analysis": analysis})
        assert response.status_code == 200, (module, response.text)
    improvement = client.post("/api/improvement/generate", json={"document_id": doc_id, "analysis": analysis})
    assert improvement.status_code == 200
    assert len(improvement.json()["suggestions"]) >= 1


def test_decision_lifecycle(tmp_path):
    doc_id = upload_sample(tmp_path, "decisions.docx")
    assert client.get(f"/api/improvements/decisions/{doc_id}").json()["decisions"] == []

    saved = client.post(
        "/api/improvements/decisions",
        json={"document_id": doc_id, "suggestion_id": "s1", "decision": "accepted"},
    )
    assert saved.status_code == 200

    edited = client.post(
        "/api/improvements/decisions",
        json={"document_id": doc_id, "suggestion_id": "s1", "decision": "edited", "edited_text": "Revised text."},
    )
    assert edited.status_code == 200

    decisions = client.get(f"/api/improvements/decisions/{doc_id}").json()["decisions"]
    assert len(decisions) == 1
    assert decisions[0]["decision"] == "edited"
    assert decisions[0]["edited_text"] == "Revised text."


def test_decision_rejects_invalid_choice(tmp_path):
    doc_id = upload_sample(tmp_path, "baddecision.docx")
    response = client.post(
        "/api/improvements/decisions",
        json={"document_id": doc_id, "suggestion_id": "s1", "decision": "maybe"},
    )
    assert response.status_code == 400
    assert client.get("/api/improvements/decisions/DOC-NOPE").status_code == 404


def test_format_rejects_unknown_journal(tmp_path):
    doc_id = upload_sample(tmp_path, "badjournal.docx")
    client.post(f"/api/documents/{doc_id}/analyze", json={})
    response = client.post(f"/api/documents/{doc_id}/format", json={"journal_id": "unknown-journal"})
    assert response.status_code == 404


def test_generate_requires_analysis(tmp_path):
    doc_id = upload_sample(tmp_path, "nogen.docx")
    assert client.post(f"/api/documents/{doc_id}/generate").status_code == 404


def test_downloads_missing_before_generation(tmp_path):
    doc_id = upload_sample(tmp_path, "nodownload.docx")
    client.post(f"/api/documents/{doc_id}/analyze", json={})
    for kind in ["docx", "pdf", "report", "zip"]:
        assert client.get(f"/api/documents/{doc_id}/download/{kind}").status_code == 404


def test_full_generation_records_artifacts(tmp_path):
    doc_id = upload_sample(tmp_path, "fullgen.docx")
    client.post(f"/api/documents/{doc_id}/analyze", json={"journal_id": "nature_like"})
    generated = client.post(f"/api/documents/{doc_id}/generate", json={"journal_id": "nature_like"})
    assert generated.status_code == 200
    for kind in ["docx", "pdf", "report", "zip"]:
        response = client.get(f"/api/documents/{doc_id}/download/{kind}")
        assert response.status_code == 200, kind
        assert len(response.content) > 0
    library = client.get("/api/documents").json()["documents"]
    entry = next(item for item in library if item["document_id"] == doc_id)
    assert entry["has_generated"] is True
